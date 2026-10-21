"""Modèle Excel/Google Sheet du staff d'un événement : saisie guidée (listes déroulantes) + onglets ROLE et SYNTHÈSE à compteurs.

À importer dans Google Drive (Nouveau → Importer un fichier, puis « Ouvrir avec Google Sheets »), à partager par lien en
lecture, puis à coller dans la fiche de l'événement (onglet Staff). Le premier onglet doit rester « STAFF ».
"""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, TEAL, SOFT, BAND, GRID = "1F3A5F", "2A9D8F", "EAF2F8", "F6F9FC", "D5DDE5"
SIZES = ["XS", "S", "M", "L", "XL", "XXL"]
FIRST, LAST = 5, 64            # lignes de saisie (60 places)
thin = Side(style="thin", color=GRID)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def _fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def _title(ws, text, sub, last_col):
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = text
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = _fill(NAVY)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 36
    ws.merge_cells(f"A2:{last_col}2")
    ws["A2"] = sub
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="5B6B7B")
    ws["A2"].alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[2].height = 32


def _head(ws, row, labels):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = _fill(TEAL)
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        c.border = BOX
    ws.row_dimensions[row].height = 30


def staff_template_xlsx(event_name: str, event_date: date | None, venue: str, roles: list[dict]) -> bytes:
    wb = Workbook()
    # ------------------------------------------------------------------ STAFF (saisie)
    ws = wb.active
    ws.title = "STAFF"
    when = f"{event_date:%d/%m/%Y}" if event_date else "date à confirmer"
    _title(ws, f"ADÉQUAT PADEL TOUR  ·  STAFF {event_name.upper()}",
           f"{when}{' · ' + venue if venue else ''}   —   Merci de compléter une ligne par personne. Rôle : choisissez dans la liste "
           "ou indiquez votre préférence, nous attribuerons les rôles (voir onglet ROLE).", "E")
    _head(ws, 4, ["Nom", "Prénom", "Taille polo", "Déjeuner avant l’événement ?", "Rôle (cf onglet \"ROLE\")"])
    for r in range(FIRST, LAST + 1):
        for c in range(1, 6):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="center", indent=1)
            if r % 2 == 0:
                cell.fill = _fill(BAND)
        ws.row_dimensions[r].height = 22
    for col, w in zip("ABCDE", (22, 22, 14, 30, 34)):
        ws.column_dimensions[col].width = w
    dv_size = DataValidation(type="list", formula1='"' + ",".join(SIZES) + '"', allow_blank=True)
    dv_size.error, dv_size.errorTitle = "Choisissez une taille de la liste.", "Taille polo"
    dv_lunch = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    dv_lunch.error, dv_lunch.errorTitle = "Répondez Oui ou Non.", "Déjeuner"
    dv_role = DataValidation(type="list", formula1=f"=ROLE!$A$5:$A${4 + len(roles)}", allow_blank=True, showErrorMessage=False)
    for dv, col in ((dv_size, "C"), (dv_lunch, "D"), (dv_role, "E")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{LAST}")
    ws.conditional_formatting.add(f"D{FIRST}:D{LAST}", CellIsRule(operator="equal", formula=['"Oui"'], fill=_fill("DDF3E4")))
    ws.conditional_formatting.add(f"D{FIRST}:D{LAST}", CellIsRule(operator="equal", formula=['"Non"'], fill=_fill("FDE8E4")))
    ws.freeze_panes = f"A{FIRST}"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    # ------------------------------------------------------------------ ROLE
    wr = wb.create_sheet("ROLE")
    _title(wr, "RÔLES DU JOUR J", "Nous attribuons les rôles ; si vous avez une préférence, mentionnez-la dans la colonne « Rôle » de l’onglet STAFF.", "C")
    _head(wr, 4, ["Rôle", "Description", "Effectif visé"])
    for i, r in enumerate(roles, 5):
        wr.cell(row=i, column=1, value=r["name"]).font = Font(name="Calibri", size=11, bold=True, color=NAVY)
        wr.cell(row=i, column=2, value=r.get("description", ""))
        t = wr.cell(row=i, column=3, value=r.get("target") if r.get("target") is not None else "à définir")
        t.alignment = Alignment(horizontal="center", vertical="center")
        for c in range(1, 4):
            cell = wr.cell(row=i, column=c)
            cell.border = BOX
            cell.alignment = Alignment(vertical="center", wrap_text=True, indent=0 if c == 3 else 1,
                                       horizontal="center" if c == 3 else "left")
            if i % 2 == 0:
                cell.fill = _fill(BAND)
        wr.row_dimensions[i].height = 28
    for col, w in zip("ABC", (24, 70, 16)):
        wr.column_dimensions[col].width = w
    wr.sheet_view.showGridLines = False
    # ------------------------------------------------------------------ SYNTHÈSE (compteurs automatiques)
    wsum = wb.create_sheet("SYNTHÈSE")
    _title(wsum, f"SYNTHÈSE STAFF  ·  {event_name.upper()}", "Compteurs automatiques : ne rien saisir dans cet onglet.", "D")
    kp = [("Inscrits staff", f"=COUNTA(STAFF!A{FIRST}:A{LAST})+COUNTA(STAFF!B{FIRST}:B{LAST})-SUMPRODUCT((STAFF!A{FIRST}:A{LAST}<>\"\")*(STAFF!B{FIRST}:B{LAST}<>\"\"))"),
          ("Déjeuner : Oui", f'=COUNTIF(STAFF!D{FIRST}:D{LAST},"Oui")'),
          ("Déjeuner : Non", f'=COUNTIF(STAFF!D{FIRST}:D{LAST},"Non")'),
          ("Rôles non renseignés", f'=B4-SUMPRODUCT((STAFF!E{FIRST}:E{LAST}<>"")*((STAFF!A{FIRST}:A{LAST}<>"")+(STAFF!B{FIRST}:B{LAST}<>"")>0))')]
    for i, (label, formula) in enumerate(kp, 4):
        a = wsum.cell(row=i, column=1, value=label)
        a.font = Font(name="Calibri", size=12, bold=True, color=NAVY)
        a.fill = _fill(SOFT)
        b = wsum.cell(row=i, column=2, value=formula)
        b.font = Font(name="Calibri", size=16, bold=True, color=TEAL)
        b.alignment = Alignment(horizontal="center")
        for c in (a, b):
            c.border = BOX
        wsum.row_dimensions[i].height = 26
    wsum.cell(row=9, column=1, value="TAILLES DE POLO").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    _head(wsum, 10, ["Taille", "Nombre"])
    for i, size in enumerate(SIZES, 11):
        wsum.cell(row=i, column=1, value=size).border = BOX
        c = wsum.cell(row=i, column=2, value=f'=COUNTIF(STAFF!C{FIRST}:C{LAST},"{size}")')
        c.border, c.alignment = BOX, Alignment(horizontal="center")
    r0 = 11 + len(SIZES) + 1
    wsum.cell(row=r0, column=1, value="RÔLES : ATTRIBUÉS / VISÉS").font = Font(name="Calibri", size=11, bold=True, color=NAVY)
    _head(wsum, r0 + 1, ["Rôle", "Attribués", "Visés", "Statut"])
    for i, r in enumerate(roles, r0 + 2):
        k = i - (r0 + 2) + 5
        wsum.cell(row=i, column=1, value=f"=ROLE!A{k}")
        wsum.cell(row=i, column=2, value=f'=COUNTIF(STAFF!E{FIRST}:E{LAST},"*"&A{i}&"*")')
        wsum.cell(row=i, column=3, value=f"=ROLE!C{k}")
        wsum.cell(row=i, column=4, value=f'=IF(ISNUMBER(C{i}),IF(B{i}>=C{i},"✓ complet","manque "&(C{i}-B{i})),"—")')
        for c in range(1, 5):
            cell = wsum.cell(row=i, column=c)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="left" if c in (1, 4) else "center", indent=1 if c in (1, 4) else 0)
    rng = f"D{r0 + 2}:D{r0 + 1 + len(roles)}"
    wsum.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(D{r0 + 2},1)="✓"'], fill=_fill("DDF3E4")))
    wsum.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(D{r0 + 2},6)="manque"'], fill=_fill("FDE8E4")))
    for col, w in zip("ABCD", (30, 16, 14, 18)):
        wsum.column_dimensions[col].width = w
    wsum.sheet_view.showGridLines = False
    for sheet in wb.worksheets:      # impression : une page de large
        sheet.page_setup.orientation = "landscape" if sheet.title != "STAFF" else "portrait"
        sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
