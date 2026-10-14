"""Documents Word : étiquettes Avery L7160 (3 x 7, 63,5 x 38,1 mm) et export des équipes/poules pour l'outil de tournoi."""
from __future__ import annotations

import csv
import io
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .participants import Participant
from .tournament import Plan
from .util import norm

COLS, ROWS = 3, 7
LABEL_W, LABEL_H, GAP = 63.5, 38.1, 2.5   # mm


def _fit_pt(text: str, base: float, max_chars: int, floor: float = 8) -> float:
    return max(floor, base * min(1.0, max_chars / max(len(text), 1)))


def _no_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tbl_pr.append(borders)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def _cell_margins(cell, mm: float = 3) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(mm * 56.7)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def _para(cell, text: str, size: float, bold=False, color=None, first=False, space_after=0):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.size, run.font.bold, run.font.name = Pt(size), bold, "Calibri"
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def labels_docx(path, brand: str, event_name: str, participants: list[Participant]) -> Path:
    """Planche Word prête à imprimer sur papier Avery L7160 (21 étiquettes par page A4), triée par nom."""
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(15.1), Mm(8)
    sec.left_margin = sec.right_margin = Mm((210 - (COLS * LABEL_W + (COLS - 1) * GAP)) / 2)
    sec.header_distance = sec.footer_distance = Mm(5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(0)

    people = sorted(participants, key=lambda p: (norm(p.nom), norm(p.prenom)))
    per_page = COLS * ROWS
    pages = max(1, -(-len(people) // per_page))
    ncols = COLS * 2 - 1  # étiquette, espace, étiquette, espace, étiquette
    for page in range(pages):
        table = doc.add_table(rows=ROWS, cols=ncols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _no_borders(table)
        for c, col in enumerate(table.columns):  # la grille doit porter les largeurs (sinon Word/LibreOffice les égalisent)
            col.width = Mm(LABEL_W if c % 2 == 0 else GAP)
        for r in range(ROWS):
            row = table.rows[r]
            row.height, row.height_rule = Mm(LABEL_H), WD_ROW_HEIGHT_RULE.EXACTLY
            for c in range(ncols):
                cell = row.cells[c]
                cell.width = Mm(LABEL_W if c % 2 == 0 else GAP)
                if c % 2:
                    continue
                idx = page * per_page + r * COLS + c // 2
                _cell_margins(cell)
                if idx >= len(people):
                    continue
                p = people[idx]
                name = f"{p.prenom} {p.nom.upper()}".strip() or p.email
                _para(cell, name, _fit_pt(name, 16, 22), bold=True, color=(0x1F, 0x3A, 0x5F), first=True, space_after=3)
                if p.company:
                    _para(cell, p.company, _fit_pt(p.company, 12, 28), space_after=3)
                poule = (p.poule if p.poule.lower().startswith("poule") else f"Poule {p.poule}") if p.poule else ""
                detail = " - ".join(v for v in (p.type or "Tournoi", poule) if v)
                _para(cell, detail, 10, bold=True, space_after=2)
                _para(cell, f"{brand} - {event_name}"[:60], 7, color=(0x80, 0x80, 0x80))
        if page < pages - 1:
            doc.add_page_break()
    # supprime le paragraphe vide initial laissé par python-docx
    first = doc.paragraphs[0] if doc.paragraphs else None
    if first is not None and not first.text and first._p.getparent() is doc.element.body:
        first._p.getparent().remove(first._p)
    doc.save(str(path))
    return Path(path)


EXPORT_HEADER = ["Poule", "Équipe", "Joueur 1", "Entreprise 1", "Joueur 2", "Entreprise 2", "Niveau équipe"]


def tournament_rows(plan: Plan) -> list[list[str]]:
    rows = []
    for poule in plan.poules:
        for team in sorted(poule.teams, key=lambda t: (-t.level, norm(t.name))):
            a = team.members[0]
            b = team.members[1] if len(team.members) > 1 else None
            rows.append([poule.name, team.name, a.full_name, a.company,
                         b.full_name if b else "", b.company if b else "", f"{team.level:.1f}"])
    return rows


def tournament_csv(plan: Plan) -> bytes:
    """CSV (séparateur « ; », UTF-8 avec BOM : s'ouvre directement dans Excel) : une ligne = une équipe."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(EXPORT_HEADER)
    w.writerows(tournament_rows(plan))
    return buf.getvalue().encode("utf-8-sig")


def tournament_xlsx(plan: Plan) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Équipes"
    ws.append(EXPORT_HEADER)
    for r in tournament_rows(plan):
        ws.append(r)
    for col, width in zip("ABCDEFG", (8, 30, 24, 24, 24, 24, 14)):
        ws.column_dimensions[col].width = width
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
