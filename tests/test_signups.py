"""Inscriptions du formulaire : solo / équipe, idempotence, poules, étiquettes Word, export tournoi."""
from datetime import date
from pathlib import Path

import pytest
from docx import Document

from padel import importer, store
from padel.config import load_settings
from padel.db import init_db, make_engine
from padel.documents_docx import labels_docx, tournament_csv
from padel.participants import Participant
from padel.signups import parse_signups
from padel.tournament import plan_tournament

CSV = Path(__file__).parent.parent / "examples" / "inscriptions-angers-exemple.csv"


@pytest.fixture()
def env(tmp_path):
    eng = make_engine(f"sqlite:///{tmp_path / 't.db'}")
    init_db(eng)
    with eng.begin() as cx:
        store.upsert_event(cx, "ANGERS", date=date(2026, 10, 15))
    return eng, load_settings()


def _parts(cx):
    ev = store.get_event(cx, "ANGERS")
    return [Participant(id=str(r["id"]), prenom=r["prenom"], nom=r["nom"], email=r["email"], company=r["company"],
                        team=r["team"], level=r["level"], type=r["type"]) for r in store.list_registrations(cx, ev["id"])]


def test_parse_solo_and_team():
    rows = [{"prenom": "Ana", "nom": "Diaz", "email": "a@x.fr", "niveau": "Confirmé", "enjeux": "Réseau",
             "prenom du coequipier": "Bob", "nom du coequipier": "Ray", "niveau du coequipier": "Débutant"},
            {"prenom": "Seul", "nom": "Joueur", "email": "s@x.fr", "niveau": "Débutant"}]
    out, warns, _ = parse_signups(rows, load_settings())
    assert [len(s["players"]) for s in out] == [2, 1]
    assert out[0]["team"] == "Diaz / Ray" and out[0]["stakes"] == "Réseau"
    assert out[0]["players"][1]["level"] == "Débutant" and out[1]["team"] == ""


def test_import_is_idempotent_and_keeps_details(env):
    eng, s = env
    with eng.begin() as cx:
        r1 = importer.import_signups(cx, CSV.read_text(encoding="utf-8-sig"), "ANGERS", s)
        r2 = importer.import_signups(cx, CSV.read_text(encoding="utf-8-sig"), "ANGERS", s)
        regs = store.list_registrations(cx, store.get_event(cx, "ANGERS")["id"])
    assert r1.created == 15 and r2.created == 0
    assert len(regs) == 15 and sum(1 for r in regs if not r["team"]) == 3
    assert all(r["email"] and r["phone"] and r["stakes"] for r in regs if r["prenom"] in ("Camille", "Hugo"))
    assert any("niveau du coéquipier" in w for w in r1.warnings)


def test_dry_run_writes_nothing(env):
    eng, s = env
    with eng.begin() as cx:
        importer.import_signups(cx, CSV.read_text(encoding="utf-8-sig"), "ANGERS", s, dry_run=True)
        assert store.list_registrations(cx) == []


def test_pools_labels_and_export(env, tmp_path):
    eng, s = env
    with eng.begin() as cx:
        importer.import_signups(cx, CSV.read_text(encoding="utf-8-sig"), "ANGERS", s)
        parts = _parts(cx)
        plan = plan_tournament(parts, s, date(2026, 10, 15))
        mapping = {int(m.id): po.name for po in plan.poules for t in po.teams for m in t.members}
        assert store.set_poules(cx, store.get_event(cx, "ANGERS")["id"], mapping) == 15
    assert len(plan.teams) == 8 and len(plan.poules) == 2
    doc = Document(str(labels_docx(tmp_path / "l.docx", "Brand", "ANGERS", parts)))
    text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "Camille DURAND" in text and len(doc.tables) == 1
    lines = tournament_csv(plan).decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Poule;Équipe") and len(lines) == 9


def test_label_pages(env, tmp_path):
    people = [Participant(id=str(i), prenom=f"P{i}", nom=f"N{i:02d}") for i in range(25)]
    doc = Document(str(labels_docx(tmp_path / "l.docx", "B", "E", people)))
    assert len(doc.tables) == 2  # 21 étiquettes par page


# ----------------------------------------------------------------------------- Typeform en direct
class FakeTF:
    TITLES = {"f1": "Prénom", "f2": "Nom", "f3": "Votre adresse e-mail", "f4": "Téléphone", "f5": "Entreprise",
              "f6": "Votre niveau de padel", "f7": "Vos enjeux pour cette journée", "f8": "Prénom de votre coéquipier",
              "f9": "Nom de votre coéquipier", "f10": "E-mail de votre coéquipier"}

    def __init__(self, responses):
        self.responses = responses

    def get_form(self, form_id):
        return {"fields": [{"id": k, "title": v, "type": "short_text"} for k, v in self.TITLES.items()]}

    def iter_responses(self, form_id):
        yield from self.responses


def _resp(token, when, **ans):
    names = {"prenom": "f1", "nom": "f2", "mail": "f3", "tel": "f4", "ent": "f5", "niv": "f6", "enj": "f7",
             "pp": "f8", "pn": "f9", "pm": "f10"}
    return {"token": token, "submitted_at": when,
            "answers": [{"field": {"id": names[k]}, "type": "text", "text": v} for k, v in ans.items()]}


RESPONSES = [
    _resp("t1", "2026-10-01T08:00:00Z", prenom="Alice", nom="Roy", mail="alice@ex.fr", tel="0600000001", ent="Acme",
          niv="Confirmé", enj="Réseau", pp="Bruno", pn="Lam", pm="bruno@ex.fr"),
    _resp("t2", "2026-10-02T09:00:00Z", prenom="Chloé", nom="Vidal", mail="chloe@ex.fr", ent="Beta", niv="Débutant"),
]


def test_typeform_sync_creates_team_and_solo_and_is_idempotent(env):
    eng, s = env
    with eng.begin() as cx:
        store.update_event(cx, store.get_event(cx, "ANGERS")["id"], typeform_id="FORM1")
        sim = importer.sync_typeform(cx, "ANGERS", FakeTF(RESPONSES), s, dry_run=True)
        assert store.list_registrations(cx) == [] and sim.created == 3
        r1 = importer.sync_typeform(cx, "ANGERS", FakeTF(RESPONSES), s)
        r2 = importer.sync_typeform(cx, "ANGERS", FakeTF(RESPONSES), s)
        regs = store.list_registrations(cx)
        k = store.participation(regs, 60)
    assert r1.created == 3 and r2.created == 0 and len(regs) == 3
    assert k == {"participants": 3, "teams": 1, "solos": 1, "teams_total": 2, "odd": True, "places_left": 57}
    lam = next(r for r in regs if r["nom"] == "Lam")
    assert lam["team"] == "Lam / Roy" and lam["stakes"] == "Réseau" and lam["level"] == "Confirmé"


def test_engine_syncs_typeform_only_with_a_token(env, monkeypatch):
    from padel import engine
    eng, s = env
    with eng.begin() as cx:
        store.update_event(cx, store.get_event(cx, "ANGERS")["id"], typeform_id="FORM1")
    engine.run(eng, s, today=date(2026, 10, 3))
    with eng.connect() as cx:
        assert store.list_registrations(cx) == []          # pas de jeton : rien n'est simulé
    s.env["TYPEFORM_TOKEN"] = "tfp_test"
    monkeypatch.setattr(type(s), "typeform_client", lambda self: FakeTF(RESPONSES))
    res = engine.run(eng, s, today=date(2026, 10, 3))
    with eng.connect() as cx:
        assert len(store.list_registrations(cx)) == 3
    assert res.errors == 0


def test_headers_are_read_from_all_rows_and_diagnosed():
    """Une 1re réponse vide ne doit pas masquer les questions ; si rien n'est reconnu, les intitulés sont listés."""
    rows = [{"submitted_at": "2026-10-01 08:00:00"},
            {"submitted_at": "2026-10-02 08:00:00", "prenom": "Ana", "nom": "Diaz", "votre e-mail": "a@x.fr"}]
    out, warns, _ = parse_signups(rows, load_settings())
    assert len(out) == 1 and out[0]["players"][0]["email"] == "a@x.fr" and not warns
    _, warns, _ = parse_signups([{"quel est votre identite": "x"}], load_settings())
    assert any(w.startswith("Questions/colonnes lues") and "identite" in w for w in warns)


def test_contact_block_and_hidden_fields_are_read():
    from padel.signups import typeform_diagnostic, typeform_rows
    form = {"fields": [
        {"id": "c", "type": "contact_info", "title": "Vos coordonnées", "properties": {"fields": [
            {"id": "c1", "type": "short_text", "ref": "first_name", "title": ""},
            {"id": "c2", "type": "short_text", "ref": "last_name", "title": ""},
            {"id": "c3", "type": "email", "ref": "email", "title": ""},
            {"id": "c4", "type": "phone_number", "ref": "phone_number", "title": ""}]}},
        {"id": "q", "type": "short_text", "title": "Quels sont vos enjeux ?"}]}
    resp = [{"submitted_at": "2026-10-01T08:00:00Z", "hidden": {"niveau": "Confirmé"},
             "answers": [{"field": {"id": "c1"}, "type": "text", "text": "Ana"}, {"field": {"id": "c2"}, "type": "text", "text": "Diaz"},
                         {"field": {"id": "c3"}, "type": "email", "email": "a@x.fr"},
                         {"field": {"id": "c4"}, "type": "phone_number", "phone_number": "+33600000000"},
                         {"field": {"id": "q"}, "type": "text", "text": "Réseau"}]}]
    out, warns, _ = parse_signups(typeform_rows(form, resp), load_settings())
    p = out[0]["players"][0]
    assert (p["prenom"], p["nom"], p["email"], p["phone"], p["level"]) == ("Ana", "Diaz", "a@x.fr", "+33600000000", "Confirmé")
    assert out[0]["stakes"] == "Réseau" and not warns
    diag = "\n".join(typeform_diagnostic(form, resp))
    assert "Ana" not in diag and "a@x.fr" not in diag and "first_name" in diag


def test_second_contact_block_is_the_partner():
    """Deux blocs de coordonnées aux intitulés identiques : le 2e (collaborateur) forme une équipe avec le 1er."""
    from padel.signups import typeform_rows

    def block(i, title):
        return {"id": f"b{i}", "type": "contact_info", "title": title, "properties": {"fields": [
            {"id": f"p{i}", "type": "short_text", "title": "Prénom"}, {"id": f"n{i}", "type": "short_text", "title": "Nom"},
            {"id": f"e{i}", "type": "email", "title": "E-mail"}]}}
    for second_title in ("Coordonnées de votre collaborateur", "Remplir vos coordonnées"):  # avec ou sans mot-clé
        form = {"fields": [block(1, "Remplir vos coordonnées"), block(2, second_title)]}
        ans = lambda fid, t, v: {"field": {"id": fid}, "type": t, t: v}  # noqa: E731
        resp = [{"submitted_at": "2026-10-01T08:00:00Z", "answers": [
            ans("p1", "text", "Ana"), ans("n1", "text", "Diaz"), ans("e1", "email", "a@x.fr"),
            ans("p2", "text", "Bob"), ans("n2", "text", "Ray"), ans("e2", "email", "b@x.fr")]}]
        out, warns, _ = parse_signups(typeform_rows(form, resp), load_settings())
        assert len(out) == 1 and len(out[0]["players"]) == 2, second_title
        assert out[0]["team"] == "Diaz / Ray" and out[0]["players"][1]["email"] == "b@x.fr"
