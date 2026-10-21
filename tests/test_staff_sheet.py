"""Staff : lecture d'un Google Sheet partagé par lien, synchronisation, compteurs, modèle."""
from datetime import date

import pytest

from padel import staff_sheet, store
from padel.config import load_settings
from padel.db import init_db, make_engine
from padel.staff_template import staff_template_xlsx

CSV = ('ADÉQUAT PADEL TOUR · STAFF BORDEAUX,,,,\n"19/11/2026 — Merci",,,,\n,,,,\n'
       'Nom,Prénom,Taille polo,Déjeuner avant l’événement ?,"Rôle (cf onglet ""ROLE"")"\n'
       'Martin,Julie,M,Oui,Field lead\nDurand,Paul,l,non,Reception\nRoy,Ana,XL,oui,Field lead (préférence)\n,,,,\nVidal,Sam,?,,\n')


@pytest.fixture()
def env(tmp_path):
    eng = make_engine(f"sqlite:///{tmp_path / 't.db'}")
    init_db(eng)
    with eng.begin() as cx:
        eid, _ = store.upsert_event(cx, "BORDEAUX", date=date(2026, 11, 19))
    return eng, eid


def test_url_parsing_and_export_url():
    u = "https://docs.google.com/spreadsheets/d/19f4Addk_BATn2_2ftuLQNAW2-QmOXlnGhwN6dvnK-64/edit?usp=sharing"
    assert staff_sheet.parse_sheet_url(u) == ("19f4Addk_BATn2_2ftuLQNAW2-QmOXlnGhwN6dvnK-64", None)
    assert staff_sheet.export_url(u + "#gid=123").endswith("export?format=csv&gid=123")
    with pytest.raises(staff_sheet.SheetError):
        staff_sheet.parse_sheet_url("https://exemple.fr/x")


def test_parse_skips_banner_and_normalizes():
    members, warns = staff_sheet.parse_staff_csv(CSV)
    assert [m["nom"] for m in members] == ["Martin", "Durand", "Roy", "Vidal"]
    assert members[1]["shirt_size"] == "L" and members[1]["repas"] == "Non" and members[0]["repas"] == "Oui"
    assert any("taille" in w for w in warns)


def test_sync_adds_updates_removes_and_keeps_manual(env):
    eng, eid = env
    with eng.begin() as cx:
        cx.execute(staff_sheet.insert(staff_sheet.staff).values(event_id=eid, nom="Manuel", prenom="Marc", source=""))
        r = staff_sheet.sync_staff(cx, eid, CSV, dry_run=True)
        assert r["added"] == 4 and store.list_staff(cx, eid)[0]["nom"] == "Manuel" and len(store.list_staff(cx, eid)) == 1
        staff_sheet.sync_staff(cx, eid, CSV)
        assert {x["nom"] for x in store.list_staff(cx, eid)} == {"Manuel", "Martin", "Durand", "Roy", "Vidal"}
        again = staff_sheet.sync_staff(cx, eid, CSV)
        assert (again["added"], again["updated"], again["removed"]) == (0, 0, 0)
        newer = CSV.replace("Durand,Paul,l,non,Reception", "Durand,Paul,L,oui,Shuttle").replace("Vidal,Sam,?,,\n", "")
        r = staff_sheet.sync_staff(cx, eid, newer)
        assert (r["updated"], r["removed"]) == (1, 1)
        rows = {x["nom"]: x for x in store.list_staff(cx, eid)}
    assert rows["Durand"]["role"] == "Shuttle" and "Vidal" not in rows and "Manuel" in rows


def test_summary_counts_roles_lunch_sizes():
    roles = load_settings().get("staff", "roles")
    members, _ = staff_sheet.parse_staff_csv(CSV)
    sm = store.staff_summary(members, roles)
    assert sm["total"] == 4 and sm["lunch"] == 2 and sm["by_role"]["Field lead"] == 2 and sm["by_role"]["Reception"] == 1
    assert sm["unassigned"] == 1 and sm["sizes"]["M"] == 1


def test_html_response_means_not_shared(monkeypatch):
    class R:
        status_code, url, headers, content = 200, "https://accounts.google.com/ServiceLogin", {"content-type": "text/html"}, b""

    class S:
        def get(self, *a, **k):
            return R()
    with pytest.raises(staff_sheet.SheetError, match="Tous les utilisateurs"):
        staff_sheet.fetch_csv("https://docs.google.com/spreadsheets/d/abc/edit", S())


def test_template_has_three_tabs_and_headers():
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(staff_template_xlsx("Bordeaux", date(2026, 11, 19), "Salle", load_settings().get("staff", "roles"))))
    assert wb.sheetnames == ["STAFF", "ROLE", "SYNTHÈSE"]
    members, _ = staff_sheet.parse_staff_csv("\n".join(",".join(str(c or "") for c in row) for row in wb["STAFF"].iter_rows(max_row=6, values_only=True)))
    assert members == []                       # modèle vide : en-têtes reconnues, aucune personne
